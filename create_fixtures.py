import os
from docx import Document
from pptx import Presentation

# Ensure directory exists
os.makedirs('tests/fixtures', exist_ok=True)

# 1. Create sample.docx
doc = Document()
doc.add_heading('Section 1: General Overview', 1)
doc.add_paragraph('This document describes the RAG system for GKN Aerospace.')
doc.add_heading('1.1 Security Protocols', 2)
doc.add_paragraph('All data must be ITAR/EAR compliant.')
doc.add_paragraph('Authorized access only.')
doc.save('tests/fixtures/sample.docx')

# 2. Create sample.pptx
prs = Presentation()

# Slide 1
slide1 = prs.slides.add_slide(prs.slide_layouts[0])
slide1.shapes.title.text = "GKN RAG Pipeline"
slide1.placeholders[1].text = "Introduction to the retrieval system."

# Slide 2
slide2 = prs.slides.add_slide(prs.slide_layouts[1])
slide2.shapes.title.text = "Features"
body = slide2.shapes.placeholders[1].text_frame
body.text = "Hierarchical parsing"
p = body.add_paragraph()
p.text = "Vector search"
p.level = 1

# Add speaker notes to slide 2
notes_slide = slide2.notes_slide
notes_slide.notes_text_frame.text = "Note: Ensure pgvector is installed on the host."

# Slide 3
slide3 = prs.slides.add_slide(prs.slide_layouts[1])
slide3.shapes.title.text = "Deployment"
slide3.shapes.placeholders[1].text = "Use Docker Compose for all environments."

prs.save('tests/fixtures/sample.pptx')

print("Created binary fixtures successfully.")
